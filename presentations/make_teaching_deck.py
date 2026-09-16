#!/usr/bin/env python3
"""Teaching deck: LFS from the proposal to the experiments, one visual per idea.
Layout: navy title band, optional grey method line, one fitted image, one takeaway box, plain-language notes.
Checks after saving: word caps, banned words, every image present, takeaway and method length."""
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[2]
F = ROOT / "docs" / "study" / "figs"
PF = ROOT / "paper" / "figs"
IF = ROOT / "paper" / "iclr2027" / "figs"
OUT = [ROOT / "paper" / "slides_pptx" / "Waiz_Khan_LFS_Teaching_Deck_2026-09-10.pptx",
       ROOT / "deliverables" / "2026-09-10" / "Waiz_Khan_LFS_Teaching_Deck_2026-09-10.pptx"]

NAVY = RGBColor(0x0B, 0x2A, 0x4A); WHITE = RGBColor(0xFF, 0xFF, 0xFF); GREY = RGBColor(0x66, 0x66, 0x66)
BOX = RGBColor(0xEE, 0xF3, 0xF9); INK = RGBColor(0x11, 0x11, 0x11)
W, H = 13.333, 7.5
BANNED = ["fingerprint", "toy", "battery", "tripwire", "u-shaped", "destroyed", "baseline", "beats", "first of its kind",
          "campaign", "cohort", "walk-forward", "headline", "deflat", "ladder", "harness", "concept direction",
          "confounds removed", "meaning-weighted", "language-agnostic", "ranked first", "scored best", "—"]
MAX_TAKEAWAY_CHARS = 250   # two lines at 14 pt across the box
MAX_METHOD_CHARS = 210     # two lines at 11 pt

SLIDES = []
def slide(title, image=None, takeaway=None, notes="", bullets=None, method=None):
    SLIDES.append(dict(title=title, image=image, takeaway=takeaway, notes=notes, bullets=bullets, method=method))

PL = F / "pipeline"
# ------------------------------------------------------------------ content
slide("Language-Factor Share", None, None,
      "This deck explains one measurement, the Language-Factor Share, from the problem it was built for to the experiments that matter. Each slide has one picture and one sentence. The notes under each slide are the words to say.",
      bullets=["What it measures, what we found, and where it stops being useful",
               "Waiz Khan, Johns Hopkins University, September 2026",
               "Aim 1 of the Koehn and Murray NSF project on multilingual language models"])

slide("The problem", PL / "c1_problem.png",
      "One model, one reading test, ten languages: 0.95 in English, 0.48 in Amharic. To close the gap we first have to see what the model does inside.",
      "A language model learns to talk by reading a huge pile of text, most of it English. Give the same reading test in ten languages and the scores fall from 0.95 in English to 0.48 in Amharic. These are real zero-shot Belebele scores for Qwen3-8B. The professors want the model to be equally good in every language. But the model is a box of numbers nobody wrote by hand, so nobody can point at the broken part. Aim 1 of the proposal is the step before fixing: look inside and understand.")

slide("What the proposal asks (Aim 1)", PL / "c2_proposal.png",
      "The proposal's picture: take in the word, think about the content, pick the output word. Its open question: how strong is the language part at each stage?",
      "Section 3.1.1 of the proposal describes three stages. The first takes in the word, the last picks the output word, and both should depend on the language. The middle stage does the thinking about content and, the professors argue, should be the same for every language. Then they write the sentence this whole project answers: it is not clear how strong the language component in the representation is, or should be. They also ask whether the spaces of different languages can be joined by a simple map, as old word vectors could.")

slide("Two possible pictures of the inside", PL / "c3_two_pictures.png",
      "Either every language meets in one shared space, or there is a big English space with small copies. The fix is different in each case, so Aim 1 asks for a measurement.",
      "Here are the two possibilities. Picture A: one shared space, where the same sentence lands in the same spot whatever the language. Color is the language, shape is the sentence. Picture B: a big English space and a small copy for every other language, and copies are never as good as the original. If A is true, small fixes can work. If B is true, only big fixes can. The proposal bets on A for the middle stage but does not know, and it asks for a measurement that can tell.")

slide("Step 1: from a sentence to one vector per layer", PL / "pl1_sentence_to_vectors.png",
      "Tokens become vectors, the transformer blocks rewrite them layer by layer, and the mean over tokens gives one vector per sentence per layer.",
      "A sentence is split into tokens, and each token starts as a vector of 4,096 numbers. The transformer blocks rewrite these vectors one layer at a time; layer 0 is the embedding output and layer B the last block. At every layer we take the hidden states, one row per token, and average them over the tokens in full precision. That gives one vector per sentence at every layer. Nothing is generated; we only read the states.",
      method="Notation: tokens x_1..x_T; hidden states h_t^(l) in R^D; pooled vector h̄^(l) = (1/T) Σ_t h_t^(l), in fp32, for every layer l = 0..B.")

slide("Step 2: the language-by-sentence grid", PL / "pl2_grid.png",
      "The same 300 sentences in 128 languages, each run once through the frozen model. At each layer the pooled vectors fill a grid: languages down, sentences across.",
      "The data are 300 news sentences professionally translated into 128 languages, the NTREX corpus. Every text is run once through the frozen model and pooled as in step 1. Arranged by language and sentence, the vectors form a grid: down a column is the same sentence in different languages, across a row is the same language telling different sentences. There is one grid per layer, and each cell holds a 4,096-dimensional vector.",
      method="Notation: s_lc = sentence c in language l; grid at layer l is X^(l) in R^(L×N×D), L = 128, N = 300, D = 4,096; cell (l, c) holds h̄_lc^(l).")

slide("Step 3: put every coordinate on the same scale", PL / "pl3_standardize.png",
      "Each of the 4,096 coordinates is standardized over the whole grid, jointly across languages, so that no single loud coordinate dominates the sums.",
      "Before any arithmetic, every coordinate is standardized: subtract its mean over the whole grid and divide by its standard deviation over the whole grid. This is done jointly across all languages, never per language, because per-language standardization would remove exactly the language differences we want to measure. It matters: in Qwen3-8B one coordinate carries 92.8 percent of the raw spread at layer 19, and without this step the share would read 0.347 instead of 0.608.",
      method="Notation: z_lcd = (h_lcd − h̄_··d) / s_d, with the mean h̄_··d and standard deviation s_d taken over all L·N cells of the grid.")

slide("Step 4: three kinds of mean", PL / "pl4_means.png",
      "Row averages give one vector per language, column averages one vector per sentence, and the average of all cells the grand mean.",
      "From the standardized grid we take three kinds of average. Averaging across a row gives the language mean, what that language looks like on average. Averaging down a column gives the sentence mean, what that sentence looks like across languages. Averaging every cell gives the grand mean. Everything that follows is built from how far these means sit from each other.",
      method="Notation: z̄_l = (1/N) Σ_c z_lc (language mean); z̄_c = (1/L) Σ_l z_lc (sentence mean); z̄ = (1/LN) Σ_lc z_lc (grand mean).")

slide("Step 5: split the total spread into three parts", PL / "pl5_sums_of_squares.png",
      "The language part measures how far the language means sit from the grand mean; the sentence part does the same for sentence means; the residual is what is left.",
      "The total spread of the grid, the squared distance of every cell from the grand mean, splits exactly into three parts. The language part is how far the language means sit from the grand mean, counted N times because each language mean stands for N cells. The sentence part is the same for sentence means, counted L times. The residual is what is specific to one language-sentence pair: interaction plus noise. On a balanced grid the three add up exactly to the total.",
      method="Notation: SS_lang = N Σ_l ||z̄_l − z̄||²;  SS_con = L Σ_c ||z̄_c − z̄||²;  SS_res = Σ_lc ||z_lc − z̄_l − z̄_c + z̄||²;  total = sum of the three.")

slide("Step 6: the Language-Factor Share", PL / "pl6_share.png",
      "LFS is the language part as a share of the two systematic parts. For Qwen3-8B at layer 19: 55.2 million against 35.6 million, so 0.608.",
      "The Language-Factor Share is the language part divided by the language part plus the sentence part. The residual is left out on purpose, so that a noisy model does not look language-neutral just because its residual is large. The share is 1 when only language separates the cells and 0 when only the sentence does. Under pure noise it is not zero but 0.298 on this grid, because averages of random numbers still differ by chance; every real value is compared with that number. For Qwen3-8B at layer 19 the real numbers are 55.2 million against 35.6 million, so 0.608.",
      method="Notation: LFS^(l) = SS_lang / (SS_lang + SS_con);  expected value under noise (L − 1)/(L + N − 2) = 127/426 = 0.298.")

slide("Step 7: repeat at every layer, read three numbers", PL / "pl7_profile.png",
      "One share per layer gives the depth profile. We summarize it by the minimum layer, the minimum LFS, and the dip depth (layer-0 value minus the minimum).",
      "The whole computation is repeated at every layer, giving B plus one numbers, one per layer. Plotted in order they form the depth profile. Three numbers summarize it: the layer where the share is smallest, the share at that layer, and the dip depth, which is the layer-0 value minus the minimum. For Qwen3-8B the minimum is at layer 19, the minimum share is 0.608, and the dip depth is 0.336. The dotted line is the noise value; every point sits far above it.",
      method="Notation: l* = argmin_l LFS^(l);  minimum LFS = LFS^(l*);  dip depth = LFS^(0) − LFS^(l*).")

slide("Step 8: the same steps on a two-by-three grid", PL / "pl8_worked_example.png",
      "With one number per cell and six cells you can do every step by hand: language part 13.5, sentence part 64, residual 0, share 0.17.",
      "To see the arithmetic, shrink the grid to two languages and three sentences with one number per cell. The language means are 6 and 9, the sentence means 3.5, 7.5 and 11.5, the grand mean 7.5. The language part is three times the two squared distances 2.25 and 2.25, which is 13.5. The sentence part is two times 16 plus 0 plus 16, which is 64. The residual is zero because this grid is exactly additive, and the parts add to the total 77.5. The share is 13.5 over 77.5, which is 0.17. The real computation is this, 4,096 lines at a time.",
      method="Every symbol from steps 4 to 6 evaluated on the six-cell grid; the noise value here would be (L − 1)/(L + N − 2) = 1/3.")

slide("How to read the curve", F / "steps" / "s10_reading_curve.png",
      "At the input the vectors are still words, so language rules. In the middle the same sentence lands close in every language. At the output the model must answer in your language.",
      "At the input the vectors are mostly the words themselves, so language decides where they land. In the middle the model has worked out what is being said, so the same sentence lands close together in every language and language matters least. At the output the model must produce a word in your language, so language rules again. But 0.61 is far above 0.30, so language never fully leaves: the shared space exists, with a language leftover we measure later.")

slide("Experiment 1: 17 models, every one dips and recovers", F / "teaching" / "tf_dip_depth_bar.png",
      "All 17 primary models dip in the middle and come back up. The dip depth ranges from 0.04 (BLOOM) to 0.34 (Qwen3-8B): nine-fold.",
      "We ran the whole measurement on 17 open base models from 11 families. Every one shows the same shape: high at the input, an interior minimum, high at the output. What differs is how deep the dip goes, from 0.04 for BLOOM to 0.34 for Qwen3-8B. A second set of 19 models, measured with a variance-component version of the estimator called LFS-VC, shows the same shape in 19 of 19. Colors are model families.",
      method="Method: depth profile on the 128×300 NTREX grid for 17 base models; direct sums of squares; dip depth = LFS^(0) − LFS^(l*).")

slide("Experiment 2: does size decide the dip?", PF / "fig7_mix_vs_scale.png",
      "Within one family bigger models dip deeper, but families differ far more than sizes do. Qwen3-1.7B dips deeper than most 7B models from other labs.",
      "Within the Qwen3 family the dip deepens with size: 0.206, 0.213, 0.314, 0.336. But across families at the same size the differences are much larger. Qwen3-1.7B dips deeper than every 7B model from another lab except Llama-3.1-8B, and BLOOM, the only deliberately balanced 46-language model here, barely dips at either size. Salamandra dips more at 2B than at 7B. So dip depth is set by how a model was trained more than by how big it is. We do not separate data, tokenizer, and curriculum.",
      method="Method: dip depth against parameter count, colored by model family; descriptive comparison within and across families.")

slide("Experiment 3: is the shared middle English?", F / "teaching" / "tf_hub_wins.png",
      "For each language we asked: is it predicted better from English, or from the average of the other languages? The average wins for 113 to 124 of 127.",
      "Some papers say models think in English. We tested it at the dip layer. For each of the 127 non-English languages we fit two predictors of its vectors: the English vectors of the same sentences, and the average of all the other languages' vectors. Each predictor is scored on sentences it never saw. The average wins for 113 to 124 of 127 languages depending on the model, by about 0.08 R-squared. English is the best single donor only for high-resource targets. The shared middle is multilingual, not English.",
      method="Method: per target language, ridge regression of its dip-layer states on English states vs on the average of the other languages; held-out R²; count of languages where the average wins.")

slide("Experiment 4: what kind of difference is left?", F / "story" / "story5_shift.png",
      "In the middle, the Arabic cloud is the English cloud moved sideways. Subtract the move and they overlap.",
      "At the dip layer the language leftover is real but simple. Picture the English vectors as a cloud and the Arabic vectors as another cloud. The Arabic cloud is the English cloud slid sideways, with a little stretch. We tested increasingly complicated moves, slide, stretch, turn, skew, bend, and credited each only with the error it removed on held-out sentences. The slide plus the stretch explains 63 to 97 percent of the gap. Turning adds under 1 percent.",
      method="Method: nested maps from each language into a shared reference space: offset (M1), scale (M2), rotation (M3), general linear (M4), nonlinear (M5); credit = held-out error reduction.")

slide("Experiment 4, measured on 14 models", PF / "fig15_budget_composition.png",
      "Offset plus scale explains 63 to 97 percent of the cross-language misalignment in every model; rotation never exceeds 0.7 percent.",
      "Here is the same result for all 14 models with saved vectors. Each bar is the held-out misalignment between languages at the dip layer, split into what each move explains: the sideways shift in dark, the stretch in light, then rotation, general linear, and nonlinear. The shift and stretch dominate everywhere. A document-purged re-check on Qwen3-1.7B lowered the nonlinear share further, so the nonlinear numbers are upper bounds. This is the proposal's mapping question answered: the map is simple.",
      method="Method: same nested maps at the dip layer on the 14 models with 1,500-sentence dumps at the common rank 64; bars are held-out misalignment split by map.")

slide("Experiment 5: the hint test", F / "teaching" / "tf_hint_test.png",
      "Does a first sentence in another language help the model predict the next English sentence? We compare a matched and a mismatched context.",
      "The proposal's third task asks whether information given in one language is used in another. We took an English sentence and the sentence before it from the same news story, translated into 41 languages. We measured how surprised the model was by the English sentence with the true context, and with a context of the same language and length taken from a different story. The difference is the benefit of the hint. The mismatched control matters, because foreign text from the wrong story actually hurts.",
      method="Method: R_content = [loss(mismatched context) − loss(matched context)] / same for English; context in 41 languages; English target; 100 sentence pairs.")

slide("Experiment 5 result: 19 models, frozen protocol", IF / "fig_r1_model_level.png",
      "Models with a smaller language share at the dip get more benefit from a foreign-language hint: Spearman 0.51 after adjustment, 0.77 at the model level.",
      "A four-model pilot looked promising but could not support a claim about models in general, so we froze a protocol and ran 19 new models from 13 families: 157,700 scored sequences. Each point is one model, placed by its sentence-component share at the dip, one minus LFS-VC, and its median hint benefit. Models with a smaller language share use the hint more. After adjusting for training data, family, script, and tokenizer the correlation is 0.51 with interval 0.19 to 0.71, and it stays positive when any family is removed.",
      method="Method: Spearman between 1 − LFS-VC at the dip and median R_content over 19 models; covariate-adjusted; 2,000-replicate bootstrap; leave-one-family-out.")

slide("Experiment 6: two behaviors, same measures", F / "teaching" / "tf_two_targets.png",
      "Nothing predicts pooled test scores well after adjustment, LFS least of all. The hint benefit is a different story. The measurement is behavior-specific.",
      "Here the same three measures are tested against two behaviors. Left: pooled multiple-choice accuracy on 33 models and 70 languages, 1,740 rows after covariate filtering, adjusted for training data and the other covariates. LFS-VC has no clear association, and even the reference measure MEXA reaches only 0.24. Right: the hint benefit on 19 models, where all three measures are clearly associated. Part of the reason is the target: two benchmarks on the same cells agree on only a third of their adjusted variation, so pooled scores are a noisy criterion. We report this null result as a limit, not a defect to tune away.",
      method="Method: covariate-adjusted Spearman (training tokens, language family, script, tokenizer fertility) against pooled Belebele and INCLUDE accuracy (33 models) and against R_content (19 models).")

slide("Experiment 7: a share cannot see shrinking", PF / "fig17_ratio_hides_scale.png",
      "One training objective shrank the representations twenty-fold. LFS barely moved, retrieval scores rose, and only the raw content variance caught it.",
      "This is not hypothetical. We fine-tuned Qwen3-0.6B with a word-alignment objective from the literature. Its representation norm fell from 451 to 42 and raw content variance from 13,791 to 699, twenty times smaller. LFS moved only from 0.445 to 0.493, and the retrieval score MEXA rose by 0.09 over the control. Left: the ratio separates the shrunk and the healthy model by 0.01. Right: the raw content variance separates them twenty-fold. The shrunk model matched its control on Belebele, so this is a detector of representation change, not of harm.",
      method="Method: fine-tune Qwen3-0.6B with a word-alignment objective, three seeds, 400 steps; read LFS, raw content variance, norm, MEXA, AaR, and Belebele.")

slide("How the measurement was checked", PL / "c7_checks.png",
      "Five checks, each with its result: a second corpus, resampled sentences, sealed predictions, injected faults, and instruction tuning. All five held.",
      "Before trusting the measurement we tried to break it five ways. A second corpus: the same 17 models on FLORES-200 Wikipedia sentences, 17 of 17 dip and recover, depth ordering Spearman 0.92. Resampled sentences: 2,000 subsets of 300 from 1,500, intervals of about plus or minus 0.01 for most models and the model ordering agrees with the full grid at Spearman 0.90 or better in every draw. Sealed predictions: four predictions for two unseen families, frozen before the runs, four of four held. Injected faults: shifts and rotations are recovered; uniform shrinking is invisible to any ratio, which is why the raw parts are always printed. Instruction tuning: six pairs, dip depth moves by at most 0.01.")

slide("What this answers in the proposal", PL / "c4_answers.png",
      "How strong is the language part: large everywhere, smallest in the middle. Is the middle English: no. Is the map simple: yes, mostly a shift. Does it show in behavior: for foreign-language context, yes.",
      "Back to the questions in the proposal, each with its number. How strong is the language component, layer by layer? Large everywhere and smallest in the middle: 0.945, 0.608, 0.958 for Qwen3-8B, with the same shape in all 36 measured profiles from 33 models. Is the shared middle English? No: the multilingual average predicts 113 to 124 of 127 languages better. Is the map between languages simple? Yes: an offset plus a scale explains 63 to 97 percent of the misalignment. Does it show in behavior? Yes for foreign-language context use, Spearman 0.51; no for pooled test scores, 0.04.")

slide("The reporting rule and the limits", PL / "c5_limits.png",
      "A share cannot see shrinking, so print the raw parts beside it. And never train against the measurement; use it to choose the layer, watch the parts, and judge a fix.",
      "Two limits, found on purpose. A share is blind to size: a big bar and a tiny bar can both be 45 percent language, so we always show the raw components next to the share. And the measurement is not a training objective: across 28 training conditions on one small model, objectives that pushed LFS down moved the retrieval-tail alignment the other way, and training on LFS itself doubled the dip while accuracy fell. Its place in the next aim is design, monitoring, and evaluation: choose the layer to act on, watch the raw components while training, and judge whether a fix moved the structure and the behavior.")

slide("What Aim 2 gets from this", PL / "c6_next_aim.png",
      "Where to act: the interior layer with the smallest language share. What to act on: the per-language offset. How to check: the raw components and a behavioral test, not the share alone.",
      "The proposal's second aim is to improve multilinguality. This work hands it three things. Where to act: the interior layer with the smallest language share, layer 19 for Qwen3-8B and layer 2 for Mistral. What to act on: the per-language offset that carries most of the language leftover; removing it from saved states already deleted 97 percent of the language part in a zero-GPU pre-test. And how to check any fix: print the raw components and run the behavioral test, never the share alone.")

slide("Say it back", None, None,
      "The closing slide is the five-sentence summary to memorize. Say it slowly, one sentence per breath. Each sentence maps onto a part of the deck: the problem, the proposal's question, the measurement, the findings, and the checks with the limits. If a listener remembers only these five sentences, they have the whole project.",
      bullets=["A model learned to talk by reading mostly English text.",
               "The professors asked how much of its inside representation is about language and how much about content.",
               "I turned every sentence into a vector, put the vectors in a language-by-sentence grid, and measured how much of the spread is about language. That share is LFS: averaging, distance from the middle, squaring, dividing.",
               "High at the input, lowest in the middle, high at the output, in all 36 measured profiles from 33 models. The middle is shared and not English. Languages differ there mostly by an offset. A smaller language share goes with more use of foreign-language context.",
               "Checked on a second corpus, resampled sentences, sealed predictions, and injected faults. It cannot see uniform shrinking, and it must not be a training objective."])

# ------------------------------------------------------------------ rendering
def fit(img_path, box_w, box_h):
    with Image.open(img_path) as im: w, h = im.size
    s = min(box_w / w, box_h / h); return w * s, h * s
def add_text(sl, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor; p = tf.paragraphs[0]; p.alignment = align; r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = "Calibri"; return tb

TAKE_Y, TAKE_H = 6.05, 1.0     # takeaway box: two lines of 14 pt with margins
METHOD_Y, METHOD_H = 1.02, 0.62  # method line: two lines of 11 pt

def build():
    prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]
    for i, s in enumerate(SLIDES, 1):
        sl = prs.slides.add_slide(blank)
        band = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Inches(0.95)); band.fill.solid(); band.fill.fore_color.rgb = NAVY; band.line.fill.background()
        add_text(sl, 0.45, 0.12, W - 0.9, 0.75, s["title"], 26, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        top = 1.1
        if s.get("method"):
            add_text(sl, 0.5, METHOD_Y, W - 1.0, METHOD_H, s["method"], 11, color=GREY); top = METHOD_Y + METHOD_H + 0.05
        if s["image"]:
            p = Path(s["image"]); assert p.exists(), f"missing image {p}"
            has_take = bool(s["takeaway"]); bottom = (TAKE_Y - 0.12) if has_take else 7.05
            box_w, box_h = W - 1.0, bottom - top
            w, h = fit(p, box_w, box_h); x = (W - w) / 2; y = top + (box_h - h) / 2
            sl.shapes.add_picture(str(p), Inches(x), Inches(y), Inches(w), Inches(h))
        if s["takeaway"]:
            bx = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(TAKE_Y), Inches(W - 1.0), Inches(TAKE_H))
            bx.fill.solid(); bx.fill.fore_color.rgb = BOX; bx.line.color.rgb = RGBColor(0xC9, 0xD6, 0xE6)
            tf = bx.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = Inches(0.25); tf.margin_top = tf.margin_bottom = Inches(0.06)
            pp = tf.paragraphs[0]; pp.alignment = PP_ALIGN.LEFT; r = pp.add_run(); r.text = s["takeaway"]; r.font.size = Pt(14); r.font.color.rgb = NAVY; r.font.name = "Calibri"
        if s["bullets"]:
            y = 1.6 if s["image"] is None else 5.4
            size = 22 if i == 1 else 17
            for k, b in enumerate(s["bullets"]):
                lines = max(1, -(-len(b) // (95 if i == 1 else 118)))   # characters per line at this size and width
                hgt = 0.2 + lines * (0.42 if i == 1 else 0.32)
                add_text(sl, 0.9, y, W - 1.8, hgt, ("" if i == 1 else "• ") + b, size, bold=(i == 1 and k == 0), color=NAVY if i == 1 else INK)
                y += hgt + (0.35 if i == 1 else 0.12)
        add_text(sl, 0.45, 7.12, 8, 0.3, "Language-Factor Share, teaching deck", 10, color=GREY)
        add_text(sl, W - 1.5, 7.12, 1.1, 0.3, f"{i} / {len(SLIDES)}", 10, color=GREY, align=PP_ALIGN.RIGHT)
        sl.notes_slide.notes_text_frame.text = s["notes"]
    return prs

def check(path):
    prs = Presentation(str(path)); problems = []
    for i, (sl, s) in enumerate(zip(prs.slides, SLIDES), 1):
        texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
        vis = " ".join(texts); notes = sl.notes_slide.notes_text_frame.text
        nvis = len(vis.split()); nnotes = len(notes.split())
        if nvis > 160: problems.append(f"slide {i}: {nvis} visible words")
        if not (30 <= nnotes <= 170): problems.append(f"slide {i}: notes {nnotes} words (want 30 to 170)")
        if s["takeaway"] and len(s["takeaway"]) > MAX_TAKEAWAY_CHARS: problems.append(f"slide {i}: takeaway {len(s['takeaway'])} chars (max {MAX_TAKEAWAY_CHARS})")
        if s.get("method") and len(s["method"]) > MAX_METHOD_CHARS: problems.append(f"slide {i}: method {len(s['method'])} chars (max {MAX_METHOD_CHARS})")
        low = (vis + " " + notes).lower()
        for b in BANNED:
            if b in low: problems.append(f"slide {i}: banned word {b!r}")
        print(f"  slide {i:2d}: visible {nvis:3d} words, notes {nnotes:3d} words")
    if problems: raise AssertionError("\n".join(problems))

if __name__ == "__main__":
    prs = build()
    for out in OUT:
        out.parent.mkdir(parents=True, exist_ok=True); prs.save(str(out)); print("wrote", out)
    check(OUT[0]); print("checks passed:", len(SLIDES), "slides")
